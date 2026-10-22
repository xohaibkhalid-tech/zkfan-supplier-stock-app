"""ZK FAN Supplier, Purchase & Factory Stock System v2 (forms-based).

Flask + SQLite, mobile-first, deliberately SIMPLE screens (like the chat
forms the user filled in): clean cards, big buttons, few fields.

Data model (v2):
- parties (suppliers) with opening (previous) balances
- purchases = header (party, date, mode piece|kg, rate/kg, cash paid +
  method) + lines (piece | dhalai | set | plain | other)
- payments: kind payment | advance | purchase_payment, with method +
  discount; outstanding per the v2 formula in db.py
- stock is item-wise from purchase lines, minus usage entries
- CashBook Pro integration is HARD OFF (user order, 21 Oct 2026)

Roles: owner = full access. storekeeper = stock + usage ONLY
(no money anywhere: no balances, rates, payments, purchases).
"""
import os
import shutil
from datetime import date
from functools import wraps

from flask import (Flask, request, redirect, url_for, session, render_template,
                   flash, jsonify, send_file, abort, g)
from werkzeug.security import generate_password_hash, check_password_hash

import db
from db import (init_db, _connect, now_iso, today_iso, next_id, audit,
                party_outstanding, party_balances, total_outstanding,
                all_party_balances, all_material_stocks,
                purchase_bakaya, material_stock, get_or_create_material,
                post_stock, compute_line, sync_purchase_stock,
                clear_purchase_stock, list_methods, method_by_id,
                seed_payment_methods, purchased_items, purchased_item)
import cashbook_link

VERSION = db.VERSION  # "v3.4-forms-c"

# ---------- statement PDF (reportlab, OPTIONAL dependency) ----------
# The app must boot and serve everything else even when reportlab is
# not installed on the server: the ledger page then hides its PDF
# button and the PDF route redirects with a friendly message.
try:
    from reportlab.lib import colors as rl_colors
    from reportlab.lib.pagesizes import A4 as RL_A4
    from reportlab.lib.styles import ParagraphStyle as RL_PStyle
    from reportlab.lib.units import mm as RL_MM
    from reportlab.platypus import (Paragraph as RL_Para,
                                    SimpleDocTemplate as RL_Doc,
                                    Spacer as RL_Spacer,
                                    Table as RL_Table,
                                    TableStyle as RL_TStyle)
    HAVE_REPORTLAB = True
except Exception:  # pragma: no cover - depends on server packages
    HAVE_REPORTLAB = False

app = Flask(__name__)
app.secret_key = os.environ.get("SUPPLIER_STOCK_SECRET", "dev-only-change-me")
app.config["MAX_CONTENT_LENGTH"] = 4 * 1024 * 1024

OWNER_USER = os.environ.get("SUPPLIER_STOCK_OWNER_USER", "owner")
OWNER_PASS = os.environ.get("SUPPLIER_STOCK_OWNER_PASSWORD", "changeme123")


# ---------- startup ----------

def _ensure_owner():
    con = _connect()
    try:
        n = con.execute("SELECT COUNT(*) FROM users").fetchone()[0]
        if n == 0:
            con.execute(
                "INSERT INTO users (username, password_hash, role,"
                " must_change_pw, active, created_at)"
                " VALUES (?,?, 'owner', 1, 1, ?)",
                (OWNER_USER, generate_password_hash(OWNER_PASS), now_iso()),
            )
            con.commit()
    finally:
        con.close()


init_db()
_ensure_owner()
seed_payment_methods()


def _auto_backup():
    """Daily timestamped backup on startup (best-effort, never breaks boot)."""
    try:
        bdir = os.path.join(db.BASE_DIR, "backups")
        os.makedirs(bdir, exist_ok=True)
        today = date.today().isoformat().replace("-", "")
        if any(f.startswith(f"supplier_stock_{today}") and f.endswith(".db")
               for f in os.listdir(bdir)):
            return
        stamp = now_iso().replace(":", "").replace("-", "")[:14]
        shutil.copy2(db.DB_PATH, os.path.join(
            bdir, f"supplier_stock_{stamp}.db"))
        files = sorted(f for f in os.listdir(bdir) if f.endswith(".db"))
        for old in files[:-30]:
            os.remove(os.path.join(bdir, old))
    except Exception:
        pass


_auto_backup()


# ---------- helpers ----------

def fmt_rs(n):
    try:
        return "Rs {:,.0f}".format(float(n or 0))
    except (TypeError, ValueError):
        return "Rs 0"


def fmt_qty(n):
    try:
        v = float(n or 0)
        return ("%.3f" % v).rstrip("0").rstrip(".")
    except (TypeError, ValueError):
        return "0"


app.jinja_env.filters["rs"] = fmt_rs
app.jinja_env.filters["qty"] = fmt_qty


def fnum(v, default=0.0):
    try:
        return float(str(v).strip() or default)
    except (TypeError, ValueError):
        return default


def _safe_next():
    """Safe internal redirect target from a `next` form field / query
    arg (v3.4, return-to-ledger): only a local path starting with
    exactly one '/' is honoured — '//host' (protocol-relative) and
    '/\\host' tricks are rejected, as is anything not starting with
    '/'. Returns '' when absent/unsafe so callers fall back to their
    default redirect."""
    nxt = (request.form.get("next") or request.args.get("next")
           or "").strip()
    if nxt.startswith("/") and not nxt.startswith("//") \
            and not nxt.startswith("/\\"):
        return nxt
    return ""


def current_user():
    # TESTING MODE (temporary, user-approved 19 Oct 2026): skip login entirely.
    # Enabled ONLY when the server env has SUPPLIER_STOCK_NO_AUTH=1.
    if os.environ.get("SUPPLIER_STOCK_NO_AUTH") == "1":
        return {"id": 0, "username": OWNER_USER, "role": "owner",
                "must_change_pw": 0, "active": 1}
    # v2.2 perf: cached on flask.g — the auth decorator, the context
    # processor and the route used to run this users-table query ~3x per
    # request (each on its own connection).
    try:
        if "_cu_loaded" in g:
            return g._cu
    except RuntimeError:
        pass
    uid = session.get("uid")
    u = None
    if uid:
        con = _connect()
        try:
            u = con.execute(
                "SELECT * FROM users WHERE id=? AND active=1",
                (uid,)).fetchone()
        finally:
            con.close()
    try:
        g._cu = u
        g._cu_loaded = True
    except RuntimeError:
        pass
    return u


def login_required(f):
    @wraps(f)
    def w(*a, **kw):
        u = current_user()
        if not u:
            return redirect(url_for("login", next=request.path))
        if u["must_change_pw"] and request.endpoint != "change_password":
            return redirect(url_for("change_password"))
        return f(*a, **kw)
    return w


def owner_required(f):
    @wraps(f)
    def w(*a, **kw):
        u = current_user()
        if not u:
            return redirect(url_for("login", next=request.path))
        if u["role"] != "owner":
            abort(403)
        if u["must_change_pw"] and request.endpoint != "change_password":
            return redirect(url_for("change_password"))
        return f(*a, **kw)
    return w


@app.context_processor
def inject_user():
    u = current_user()
    return {"me": u, "is_owner": bool(u and u["role"] == "owner"),
            "version": VERSION}


def _supplier_name(sid):
    con = _connect()
    try:
        r = con.execute("SELECT name FROM suppliers WHERE id=?",
                        (sid,)).fetchone()
        return r["name"] if r else ""
    finally:
        con.close()


def _lines_of(pid, con=None):
    own = con is None
    con = con or _connect()
    try:
        return con.execute(
            "SELECT * FROM purchase_lines WHERE purchase_id=?"
            " ORDER BY sort, id", (pid,)).fetchall()
    finally:
        if own:
            con.close()


def _bulk_lines(con, purchase_ids):
    """All lines for a set of purchases in ONE query (v2.2 perf fix —
    the list page used to run _lines_of() per purchase, each on its own
    connection). Returns {purchase_id: [line rows]}."""
    ids = list(purchase_ids)
    if not ids:
        return {}
    ph = ",".join("?" * len(ids))
    out = {}
    for ln in con.execute(
            f"SELECT * FROM purchase_lines WHERE purchase_id IN ({ph})"
            " ORDER BY sort, id", ids).fetchall():
        out.setdefault(ln["purchase_id"], []).append(ln)
    return out


def _bulk_linked(con, purchase_ids):
    """Linked later payments per purchase in ONE grouped query — the
    bulk twin of db.purchase_bakaya()'s second query.
    Returns {purchase_id: SUM(amount+discount)}."""
    ids = list(purchase_ids)
    if not ids:
        return {}
    ph = ",".join("?" * len(ids))
    return {r["purchase_id"]: (r["s"] or 0) for r in con.execute(
        "SELECT purchase_id, COALESCE(SUM(amount + discount),0) AS s"
        " FROM payments WHERE purchase_id IN (" + ph + ")"
        " AND status='active' AND kind IN ('payment','advance')"
        " GROUP BY purchase_id", ids).fetchall()}


def lines_summary(lines):
    parts = []
    for ln in lines:
        if ln["kind"] == "piece":
            parts.append(f"{ln['name']} ×{fmt_qty(ln['qty'])}")
        elif ln["kind"] in ("dhalai", "set", "plain"):
            parts.append(f"{ln['name']} {fmt_qty(ln['weight_kg'])} kg")
        else:
            parts.append(f"{ln['name']} (other)")
    return " · ".join(parts)


# ---------- public ----------

@app.route("/health")
def health():
    ok = True
    try:
        con = _connect()
        con.execute("SELECT 1").fetchone()
        con.close()
    except Exception:
        ok = False
    return jsonify({"app": "zkfan-supplier-stock", "version": VERSION,
                    "ts": now_iso(), "db": "ok" if ok else "error",
                    "cashbook": "off", "cashbook_direct": False})


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        un = (request.form.get("username") or "").strip()
        pw = request.form.get("password") or ""
        con = _connect()
        try:
            u = con.execute(
                "SELECT * FROM users WHERE username=? AND active=1",
                (un,)).fetchone()
        finally:
            con.close()
        if u and check_password_hash(u["password_hash"], pw):
            session["uid"] = u["id"]
            audit(un, "login", "user", str(u["id"]), "")
            nxt = request.args.get("next") or url_for("dashboard")
            return redirect(nxt)
        flash("Ghalat username ya password", "err")
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/change-password", methods=["GET", "POST"])
@login_required
def change_password():
    u = current_user()
    if request.method == "POST":
        p1 = request.form.get("p1") or ""
        p2 = request.form.get("p2") or ""
        if len(p1) < 6:
            flash("Password kam az kam 6 characters ho", "err")
        elif p1 != p2:
            flash("Dono password same nahi hain", "err")
        else:
            con = _connect()
            try:
                con.execute(
                    "UPDATE users SET password_hash=?, must_change_pw=0"
                    " WHERE id=?",
                    (generate_password_hash(p1), u["id"]),
                )
                con.commit()
            finally:
                con.close()
            audit(u["username"], "change_password", "user", str(u["id"]), "")
            flash("Password change ho gaya", "ok")
            return redirect(url_for("dashboard"))
    return render_template("change_password.html",
                           force=bool(u["must_change_pw"]))


# ---------- dashboard ----------

@app.route("/")
@login_required
def dashboard():
    u = current_user()
    if u["role"] != "owner":
        return redirect(url_for("stock"))
    con = _connect()
    try:
        suppliers = con.execute(
            "SELECT * FROM suppliers WHERE active=1 ORDER BY name").fetchall()
        materials = con.execute(
            "SELECT * FROM materials WHERE active=1 ORDER BY name").fetchall()
        recent_pur = con.execute(
            "SELECT p.*, s.name AS sname FROM purchases p"
            " JOIN suppliers s ON s.id=p.supplier_id"
            " ORDER BY p.date DESC, p.id DESC LIMIT 6").fetchall()
        recent_pay = con.execute(
            "SELECT p.*, s.name AS sname FROM payments p"
            " JOIN suppliers s ON s.id=p.supplier_id"
            " WHERE p.status='active'"
            " ORDER BY COALESCE(p.date,'9999') DESC, p.id DESC LIMIT 6"
        ).fetchall()
        usage7 = con.execute(
            "SELECT m.name, m.unit, COALESCE(SUM(l.qty),0) AS used"
            " FROM stock_ledger l JOIN materials m ON m.id=l.material_id"
            " WHERE l.type='OUT' AND l.date >= date('now','-7 days')"
            " GROUP BY m.id ORDER BY used DESC").fetchall()
        # v2.2 perf: bulk aggregates on THIS connection — the old code
        # called party_balances()/material_stock() once per row, each
        # opening its own connection (~269 queries for this page).
        bals = all_party_balances(
            con, {s["id"]: (s["opening_balance"] or 0) for s in suppliers})
        stocks = all_material_stocks(
            con, {m["id"]: (m["opening_stock"] or 0) for m in materials})
    finally:
        con.close()

    sup_rows, total_due, total_adv = [], 0.0, 0.0
    total_prev, total_cur = 0.0, 0.0
    for s in suppliers:
        bal = bals[s["id"]]
        sup_rows.append({"row": s, "bal": bal, "out": bal["total"]})
        total_prev += bal["prev_remaining"]
        if bal["current_bakaya"] > 0:
            total_cur += bal["current_bakaya"]
        if bal["total"] >= 0:
            total_due += bal["total"]
        else:
            total_adv += -bal["total"]
    mat_rows, low = [], []
    for m in materials:
        st = stocks.get(m["id"], 0)
        mat_rows.append({"row": m, "stock": st})
        if m["min_level"] and st < m["min_level"]:
            low.append({"row": m, "stock": st})
    pend, _fail = cashbook_link.queue_counts()
    return render_template(
        "dashboard.html", total_due=total_due, total_adv=total_adv,
        total_prev=total_prev, total_cur=total_cur,
        suppliers=sup_rows, materials=mat_rows, low_stock=low,
        recent_pur=recent_pur, recent_pay=recent_pay, usage7=usage7,
        sync_pending=pend)


# ---------- suppliers / parties (owner) ----------

CATEGORIES = ["Dhalai", "Blades", "Canopy", "Kits", "Copper", "Wire",
              "Packages", "Bearing", "Nakka", "Pipe", "Electrical",
              "Loan", "Other"]


@app.route("/suppliers")
@owner_required
def suppliers():
    con = _connect()
    try:
        rows = con.execute(
            "SELECT * FROM suppliers ORDER BY active DESC, name").fetchall()
        bals = all_party_balances(
            con, {r["id"]: (r["opening_balance"] or 0) for r in rows})
    finally:
        con.close()
    out = [{"row": r, "bal": bals[r["id"]], "out": bals[r["id"]]["total"]}
           for r in rows]
    return render_template("suppliers.html", suppliers=out)


@app.route("/suppliers/add", methods=["GET", "POST"])
@owner_required
def supplier_add():
    u = current_user()
    if request.method == "POST":
        name = (request.form.get("name") or "").strip()
        if not name:
            flash("Party ka naam zaroori hai", "err")
            return render_template("supplier_form.html", s=None,
                                   categories=CATEGORIES)
        sid = next_id("SUP", "suppliers")
        con = _connect()
        try:
            con.execute(
                "INSERT INTO suppliers (id, name, category, phone,"
                " opening_balance, opening_date, notes, active, created_at)"
                " VALUES (?,?,?,?,?,?,?,?,?)",
                (sid, name, request.form.get("category") or "Other",
                 request.form.get("phone") or "",
                 fnum(request.form.get("opening_balance")),
                 request.form.get("opening_date") or "",
                 request.form.get("notes") or "", 1, now_iso()),
            )
            con.commit()
        finally:
            con.close()
        audit(u["username"], "add", "supplier", sid, name)
        flash(f"Party add ho gayi: {name}", "ok")
        return redirect(url_for("suppliers"))
    return render_template("supplier_form.html", s=None,
                           categories=CATEGORIES)


@app.route("/suppliers/<sid>/edit", methods=["GET", "POST"])
@owner_required
def supplier_edit(sid):
    u = current_user()
    con = _connect()
    try:
        s = con.execute("SELECT * FROM suppliers WHERE id=?",
                        (sid,)).fetchone()
    finally:
        con.close()
    if not s:
        abort(404)
    if request.method == "POST":
        name = (request.form.get("name") or "").strip()
        if not name:
            flash("Naam zaroori hai", "err")
            return render_template("supplier_form.html", s=s,
                                   categories=CATEGORIES)
        con = _connect()
        try:
            con.execute(
                "UPDATE suppliers SET name=?, category=?, phone=?,"
                " opening_balance=?, opening_date=?, notes=?, active=?"
                " WHERE id=?",
                (name, request.form.get("category") or "Other",
                 request.form.get("phone") or "",
                 fnum(request.form.get("opening_balance")),
                 request.form.get("opening_date") or "",
                 request.form.get("notes") or "",
                 1 if request.form.get("active") else 0, sid),
            )
            con.commit()
        finally:
            con.close()
        audit(u["username"], "edit", "supplier", sid, name)
        flash("Party update ho gayi", "ok")
        return redirect(url_for("suppliers"))
    return render_template("supplier_form.html", s=s, categories=CATEGORIES)


@app.route("/suppliers/<sid>/delete", methods=["POST"])
@owner_required
def supplier_delete(sid):
    u = current_user()
    con = _connect()
    try:
        s = con.execute("SELECT * FROM suppliers WHERE id=?",
                        (sid,)).fetchone()
        if not s:
            con.close()
            abort(404)
        n = con.execute(
            "SELECT COUNT(*) FROM purchases WHERE supplier_id=?", (sid,)
        ).fetchone()[0] + con.execute(
            "SELECT COUNT(*) FROM payments WHERE supplier_id=?", (sid,)
        ).fetchone()[0]
        if n:
            con.close()
            flash("Is party ki purchases/payments mojood hain — pehle woh"
                  " delete karein, ya Edit se deactivate kar dein", "err")
            return redirect(url_for("suppliers"))
        con.execute("DELETE FROM suppliers WHERE id=?", (sid,))
        con.commit()
    finally:
        con.close()
    audit(u["username"], "delete", "supplier", sid, s["name"])
    flash(f"Party delete ho gayi: {s['name']}", "ok")
    return redirect(url_for("suppliers"))


# ---------- party ledger (owner, print-clean) ----------

def _ledger_rows(s, con=None):
    """Chronological khata rows for one party with a running outstanding.
    Money math mirrors db.party_outstanding; voided rows are shown struck
    through and do not move the balance. Pass a shared connection to keep
    the whole ledger page on one connection (v2.2)."""
    own = con is None
    con = con or _connect()
    try:
        purs = con.execute(
            "SELECT * FROM purchases WHERE supplier_id=?"
            " ORDER BY date, id", (s["id"],)).fetchall()
        pays = con.execute(
            "SELECT * FROM payments WHERE supplier_id=?"
            " AND kind IN ('payment','advance')"
            " ORDER BY COALESCE(date,'9999'), id", (s["id"],)).fetchall()
        mirrors = {}
        for m in con.execute(
                "SELECT * FROM payments WHERE supplier_id=?"
                " AND kind='purchase_payment'", (s["id"],)).fetchall():
            mirrors[m["purchase_id"]] = m
        lines = {}
        for ln in con.execute(
                "SELECT pl.* FROM purchase_lines pl JOIN purchases p"
                " ON p.id=pl.purchase_id WHERE p.supplier_id=?"
                " ORDER BY pl.sort, pl.id", (s["id"],)).fetchall():
            lines.setdefault(ln["purchase_id"], []).append(ln)
    finally:
        if own:
            con.close()

    rows = []
    bal = s["opening_balance"] or 0
    rows.append({"date": s["opening_date"] or "—", "type": "Opening",
                 "ref": "—", "detail": "Previous balance",
                 "plus": bal, "minus": 0, "bal": bal, "void": False,
                 "tag": "", "ekind": "opening", "eid": s["id"]})
    events = []
    for p in purs:
        events.append((p["date"] or "", 0, "PUR", p))
    for p in pays:
        events.append((p["date"] or "9999", 1, "PAY", p))
    events.sort(key=lambda e: (e[0], e[1], e[3]["id"]))
    for _d, _seq, kind, p in events:
        if kind == "PUR":
            void = p["status"] != "active"
            amt = p["amount"] or 0
            if not void:
                bal += amt
            rows.append({
                "date": p["date"], "type": "Purchase", "ref": p["id"],
                "detail": lines_summary(lines.get(p["id"], []))
                          + (f" — {p['notes']}" if p["notes"] else ""),
                "plus": amt, "minus": 0, "bal": bal, "void": void,
                "tag": "", "ekind": "purchase", "eid": p["id"]})
            m = mirrors.get(p["id"])
            if m:
                mvoid = m["status"] != "active" or void
                if not mvoid:
                    bal -= (m["amount"] or 0)
                rows.append({
                    "date": m["date"] or p["date"], "type": "Purchase cash",
                    "ref": m["id"],
                    "detail": f"{p['id']} cash — {m['method_name'] or 'Cash'}",
                    "plus": 0, "minus": m["amount"] or 0, "bal": bal,
                    "void": mvoid, "tag": "Current",
                    "ekind": "purchase", "eid": p["id"]})
        else:
            void = p["status"] != "active"
            amt = p["amount"] or 0
            disc = p["discount"] or 0
            if not void:
                bal -= (amt + disc)
            label = "Advance" if p["kind"] == "advance" else "Payment"
            det = p["method_name"] or p["paid_from"] or ""
            if disc:
                det += f" · discount {fmt_rs(disc)}"
            if p["purchase_id"]:
                det += f" · {p['purchase_id']}"
            if p["notes"]:
                det += f" — {p['notes']}"
            if p["kind"] == "advance":
                tag = "Current"
            else:
                tag = ("Current" if (p["applies_to"] or "previous")
                       == "current" else "Previous")
            rows.append({"date": p["date"] or "—", "type": label,
                         "ref": p["id"], "detail": det.strip(" ·"),
                         "plus": 0, "minus": amt + disc, "bal": bal,
                         "void": void, "tag": tag,
                         "ekind": "payment", "eid": p["id"]})
    return rows, bal


@app.route("/suppliers/<sid>/ledger")
@owner_required
def supplier_ledger(sid):
    con = _connect()
    try:
        s = con.execute("SELECT * FROM suppliers WHERE id=?",
                        (sid,)).fetchone()
        if not s:
            abort(404)
        rows, bal = _ledger_rows(s, con)
        split = party_balances(sid, con)
    finally:
        con.close()
    return render_template("party_ledger.html", s=s, rows=rows, balance=bal,
                           outstanding=split["total"],
                           split=split, have_reportlab=HAVE_REPORTLAB)


# ---------- bakaya drill-down (owner, v3.1) ----------

@app.route("/bakaya/previous")
@owner_required
def previous_bakaya():
    """Previous Bakaya — LEVEL 1 (v3.2): ONLY party rows (name + previous
    remaining + ›). Tapping a party opens its full date-wise ledger
    (supplier_ledger) where the opening balance, purchases and payments
    are all editable. The grand total is exactly the dashboard's
    PREVIOUS BAKAYA figure (same bulk split math, same active-party
    set). Parties whose previous balance is fully paid (remaining 0)
    drop off this list."""
    con = _connect()
    try:
        suppliers = con.execute(
            "SELECT * FROM suppliers WHERE active=1 ORDER BY name").fetchall()
        bals = all_party_balances(
            con, {s["id"]: (s["opening_balance"] or 0) for s in suppliers})
    finally:
        con.close()
    rows, total = [], 0.0
    for s in suppliers:
        bal = bals[s["id"]]
        total += bal["prev_remaining"]
        if bal["prev_remaining"] > 0.009:
            rows.append({"row": s, "bal": bal})
    rows.sort(key=lambda x: -x["bal"]["prev_remaining"])
    return render_template("previous_bakaya.html", rows=rows, total=total)


@app.route("/bakaya/previous/<sid>/edit", methods=["GET", "POST"])
@owner_required
def previous_bakaya_edit(sid):
    """Edit ONE party's previous (opening) balance. Guard: the new amount
    may not drop below the payments already marked applies_to='previous'
    for this party — those payments were made against this balance, so a
    smaller balance would corrupt the split math (party_balances)."""
    u = current_user()
    con = _connect()
    try:
        s = con.execute("SELECT * FROM suppliers WHERE id=?",
                        (sid,)).fetchone()
        if not s:
            abort(404)
        bal = party_balances(sid, con)
    finally:
        con.close()
    if request.method == "POST":
        raw = (request.form.get("opening_balance") or "").strip()
        try:
            amount = float(raw.replace(",", ""))
        except ValueError:
            amount = -1
        if amount < 0:
            flash("Rakam sahi likhein — 0 ya us se zyada ho", "err")
            return render_template("previous_bakaya_form.html", s=s, bal=bal)
        if amount < bal["prev_paid"] - 0.009:
            flash(f"Ye balance {fmt_rs(bal['prev_paid'])} se kam nahi ho"
                  " sakta — is party ke previous balance par itni payments"
                  " pehle hi lag chuki hain", "err")
            return render_template("previous_bakaya_form.html", s=s, bal=bal)
        con = _connect()
        try:
            con.execute(
                "UPDATE suppliers SET opening_balance=? WHERE id=?",
                (amount, sid))
            con.commit()
        finally:
            con.close()
        audit(u["username"], "edit", "previous_balance", sid,
              f"{s['name']}: {bal['opening']} -> {amount}")
        flash(f"{s['name']} ka previous balance update ho gaya:"
              f" {fmt_rs(amount)}", "ok")
        return redirect(url_for("previous_bakaya"))
    return render_template("previous_bakaya_form.html", s=s, bal=bal)


@app.route("/bakaya/previous/<sid>/delete", methods=["POST"])
@owner_required
def previous_bakaya_delete(sid):
    """Remove ONE party's previous balance (opening_balance -> 0).
    BLOCKED while any active payment is marked against it — deleting the
    balance under live payments would corrupt the split math. Zeroing is
    enough: every display derives prev_remaining from this column via
    party_balances(), so the party's previous outstanding becomes 0
    everywhere at once (dashboard, this page, party ledger)."""
    u = current_user()
    con = _connect()
    try:
        s = con.execute("SELECT * FROM suppliers WHERE id=?",
                        (sid,)).fetchone()
        if not s:
            abort(404)
        bal = party_balances(sid, con)
        if bal["prev_paid"] > 0.009:
            flash(f"{s['name']} ka previous balance delete nahi ho sakta"
                  f" — is par {fmt_rs(bal['prev_paid'])} payments lag"
                  " chuki hain. Pehle woh payments void/delete karein,"
                  " phir delete hoga", "err")
            return redirect(url_for("previous_bakaya"))
        con.execute(
            "UPDATE suppliers SET opening_balance=0 WHERE id=?", (sid,))
        con.commit()
    finally:
        con.close()
    audit(u["username"], "delete", "previous_balance", sid,
          f"{s['name']}: {bal['opening']}")
    flash(f"{s['name']} ka previous balance delete ho gaya", "ok")
    return redirect(url_for("previous_bakaya"))


def _bakaya_purchase_rows(con):
    """Active purchases of active parties that still carry bakaya
    (amount - header cash - linked later payments > 0), newest first.
    Single source for the Current Bakaya Excel export (v3.2)."""
    purs = con.execute(
        "SELECT p.*, s.name AS sname FROM purchases p"
        " JOIN suppliers s ON s.id=p.supplier_id"
        " WHERE p.status='active' AND s.active=1"
        " ORDER BY p.date DESC, p.id DESC").fetchall()
    linked = _bulk_linked(con, [p["id"] for p in purs])
    out = []
    for p in purs:
        link = linked.get(p["id"], 0)
        bakaya = (p["amount"] or 0) - (p["cash_paid"] or 0) - link
        if bakaya <= 0.009:
            continue
        out.append({"row": p, "linked": link, "bakaya": bakaya,
                    "cash_diya": (p["cash_paid"] or 0) + link})
    return out


@app.route("/bakaya/current")
@owner_required
def current_bakaya():
    """Current Bakaya — LEVEL 1 (v3.2): ONLY party rows (name + current
    bakaya + ›). Tapping a party opens its full date-wise ledger
    (supplier_ledger) where every purchase/payment is editable. The
    grand total is exactly the dashboard's CURRENT BAKAYA figure
    (positive per-party current_bakaya values summed — the dashboard's
    own rule)."""
    con = _connect()
    try:
        suppliers = con.execute(
            "SELECT * FROM suppliers WHERE active=1 ORDER BY name"
        ).fetchall()
        bals = all_party_balances(
            con, {s["id"]: (s["opening_balance"] or 0) for s in suppliers})
    finally:
        con.close()
    rows, total = [], 0.0
    for s in suppliers:
        bal = bals[s["id"]]
        if bal["current_bakaya"] > 0:
            total += bal["current_bakaya"]
        if bal["current_bakaya"] > 0.009:
            rows.append({"row": s, "bal": bal})
    rows.sort(key=lambda x: -x["bal"]["current_bakaya"])
    return render_template("current_bakaya.html", rows=rows, total=total)


# ---------- purchases v2 (owner) ----------

def _col(lst, i):
    try:
        return lst[i]
    except (IndexError, TypeError):
        return ""


def _parse_purchase_form():
    """Read the dynamic lines form -> (header dict, [computed line dicts])."""
    f = request.form
    kinds = f.getlist("kind[]")
    names = f.getlist("lname[]")
    qtys, rates = f.getlist("qty[]"), f.getlist("rate[]")
    bqs, bws = f.getlist("body_qty[]"), f.getlist("body_wt[]")
    cqs, cws = f.getlist("cup_qty[]"), f.getlist("cup_wt[]")
    wkgs, oamts = f.getlist("weight_kg[]"), f.getlist("oamount[]")
    sets_l, setw_l = f.getlist("sets[]"), f.getlist("set_wt[]")
    rate_per_kg = fnum(f.get("rate_per_kg"))
    lines = []
    for i, kind in enumerate(kinds):
        kind = (kind or "piece").strip()
        raw = {
            "name": (_col(names, i) or "").strip(),
            "qty": fnum(_col(qtys, i)), "rate": fnum(_col(rates, i)),
            "body_qty": fnum(_col(bqs, i)), "body_wt": fnum(_col(bws, i)),
            "cup_qty": fnum(_col(cqs, i)), "cup_wt": fnum(_col(cws, i)),
            "weight_kg": fnum(_col(wkgs, i)), "amount": fnum(_col(oamts, i)),
        }
        if kind == "set":
            # blade sets travel in their own fields; map onto body_qty/body_wt
            raw["body_qty"] = fnum(_col(sets_l, i))
            raw["body_wt"] = fnum(_col(setw_l, i))
            if not raw["body_qty"] and not raw["body_wt"]:
                raw["body_qty"] = fnum(_col(bqs, i))
                raw["body_wt"] = fnum(_col(bws, i))
        if not raw["name"] and not any(
                raw[k] for k in ("qty", "body_qty", "cup_qty",
                                 "weight_kg", "amount")):
            continue  # blank row
        ln = compute_line(kind, rate_per_kg=rate_per_kg, **raw)
        lines.append(ln)
    method_id = int(fnum(f.get("cash_method_id")))
    mrow = method_by_id(method_id) if method_id else None
    header = {
        "supplier_id": f.get("supplier_id") or "",
        "date": f.get("date") or today_iso(),
        "mode": (f.get("mode") or "piece").strip(),
        "rate_per_kg": rate_per_kg,
        "notes": (f.get("notes") or "").strip(),
        "cash_paid": fnum(f.get("cash_paid")),
        "cash_method_id": method_id,
        "cash_method_name": mrow["name"] if mrow else "",
    }
    return header, lines


def _mirror_purchase_payment(pid, header, actor):
    """Keep the kind='purchase_payment' mirror row in step with the
    purchase header's cash (create / update / void)."""
    con = _connect()
    try:
        row = con.execute(
            "SELECT * FROM payments WHERE purchase_id=? AND"
            " kind='purchase_payment'", (pid,)).fetchone()
        cash = header["cash_paid"]
        if cash > 0:
            if row:
                con.execute(
                    "UPDATE payments SET date=?, amount=?, method_id=?,"
                    " method_name=?, status='active' WHERE id=?",
                    (header["date"], cash, header["cash_method_id"],
                     header["cash_method_name"], row["id"]),
                )
            else:
                payid = next_id("PAY", "payments")
                con.execute(
                    """INSERT INTO payments (id, date, supplier_id, amount,
                        status, cashbook_status, notes, created_by,
                        created_at, discount, method_id, method_name, kind,
                        purchase_id, source, external_key, applies_to)
                       VALUES (?,?,?, ?, 'active', 'na', '', ?, ?, 0, ?, ?,
                               'purchase_payment', ?, '', '', 'current')""",
                    (payid, header["date"], header["supplier_id"], cash,
                     actor, now_iso(), header["cash_method_id"],
                     header["cash_method_name"], pid),
                )
        elif row:
            con.execute("UPDATE payments SET status='voided' WHERE id=?",
                        (row["id"],))
        con.commit()
    finally:
        con.close()


def _save_purchase(pid, header, lines, actor, is_new):
    amount = round(sum(ln["amount"] for ln in lines), 2)
    total_kg = round(sum(ln["weight_kg"] for ln in lines), 3)
    con = _connect()
    try:
        if is_new:
            con.execute(
                """INSERT INTO purchases (id, date, supplier_id, model,
                    total_kg, amount, status, cashbook_status, notes,
                    created_by, created_at, mode, rate_per_kg, cash_paid,
                    cash_method_id, cash_method_name)
                   VALUES (?,?,?,'',?,?,'active','na',?,?,?,?,?,?,?,?)""",
                (pid, header["date"], header["supplier_id"], total_kg,
                 amount, header["notes"], actor, now_iso(), header["mode"],
                 header["rate_per_kg"], header["cash_paid"],
                 header["cash_method_id"], header["cash_method_name"]),
            )
        else:
            con.execute(
                """UPDATE purchases SET date=?, supplier_id=?, total_kg=?,
                    amount=?, notes=?, mode=?, rate_per_kg=?, cash_paid=?,
                    cash_method_id=?, cash_method_name=? WHERE id=?""",
                (header["date"], header["supplier_id"], total_kg, amount,
                 header["notes"], header["mode"], header["rate_per_kg"],
                 header["cash_paid"], header["cash_method_id"],
                 header["cash_method_name"], pid),
            )
            con.execute("DELETE FROM purchase_lines WHERE purchase_id=?",
                        (pid,))
        for i, ln in enumerate(lines):
            con.execute(
                """INSERT INTO purchase_lines (purchase_id, kind, name, qty,
                    rate, body_qty, body_wt, cup_qty, cup_wt, body_total_kg,
                    cup_total_kg, weight_kg, amount, sort)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (pid, ln["kind"], ln["name"], ln["qty"], ln["rate"],
                 ln["body_qty"], ln["body_wt"], ln["cup_qty"], ln["cup_wt"],
                 ln["body_total_kg"], ln["cup_total_kg"], ln["weight_kg"],
                 ln["amount"], i),
            )
        con.commit()
    finally:
        con.close()
    pur = {"id": pid, "date": header["date"]}
    sync_purchase_stock(pur, actor)
    _mirror_purchase_payment(pid, header, actor)
    return amount, total_kg


@app.route("/purchases")
@owner_required
def purchases():
    f = request.args.get("f", "all")
    q = (request.args.get("q") or "").strip()
    con = _connect()
    try:
        qsql = ("SELECT p.*, s.name AS sname FROM purchases p"
                " JOIN suppliers s ON s.id=p.supplier_id")
        if f in ("active", "voided"):
            qsql += f" WHERE p.status='{f}'"
        qsql += " ORDER BY p.date DESC, p.id DESC"
        rows = con.execute(qsql).fetchall()
        results = []
        if q:
            like = ("%" + q.lower().replace("\\", "\\\\")
                    .replace("%", "\\%").replace("_", "\\_") + "%")
            srows = con.execute(
                "SELECT pl.*, p.date AS pdate, p.status AS pstatus,"
                " p.rate_per_kg AS prate, s.name AS sname"
                " FROM purchase_lines pl"
                " JOIN purchases p ON p.id = pl.purchase_id"
                " JOIN suppliers s ON s.id = p.supplier_id"
                " WHERE lower(pl.name) LIKE ? ESCAPE '\\'"
                " ORDER BY p.date DESC, p.id DESC, pl.sort, pl.id",
                (like,)).fetchall()
            for r in srows:
                if r["kind"] == "piece":
                    qty_disp = f"{fmt_qty(r['qty'])} pcs"
                    rate_disp = fmt_rs(r["rate"]) if r["rate"] else "—"
                elif r["kind"] == "other":
                    qty_disp = "—"
                    rate_disp = "—"
                else:
                    qty_disp = f"{fmt_qty(r['weight_kg'])} kg"
                    if r["kind"] == "set" and r["body_qty"]:
                        qty_disp += f" ({fmt_qty(r['body_qty'])} sets)"
                    rate_disp = (fmt_rs(r["prate"]) + "/kg"
                                 if r["prate"] else "—")
                results.append({"row": r, "qty_disp": qty_disp,
                                "rate_disp": rate_disp})
        # v2.2 perf: lines + linked payments for the whole list in two
        # bulk queries on this connection (was: 2 fresh connections and
        # ~4 queries PER purchase — the ~5s purchases page).
        lines_map = _bulk_lines(con, [r["id"] for r in rows])
        linked = _bulk_linked(
            con, [r["id"] for r in rows if r["status"] == "active"])
    finally:
        con.close()
    out = []
    for r in rows:
        lines = lines_map.get(r["id"], [])
        bakaya = ((r["amount"] or 0) - (r["cash_paid"] or 0)
                  - linked.get(r["id"], 0)) if r["status"] == "active" else 0
        out.append({"row": r, "lines": lines,
                    "summary": lines_summary(lines),
                    "bakaya": bakaya})
    return render_template("purchases.html", rows=out, f=f, q=q,
                           results=results)


@app.route("/purchases/add", methods=["GET", "POST"])
@owner_required
def purchase_add():
    u = current_user()
    con = _connect()
    try:
        suppliers = con.execute(
            "SELECT * FROM suppliers WHERE active=1 ORDER BY name").fetchall()
    finally:
        con.close()
    methods = list_methods()
    if request.method == "POST":
        header, lines = _parse_purchase_form()
        if not header["supplier_id"] or not lines:
            flash("Party aur kam az kam ek line zaroori hai", "err")
            return render_template("purchase_form.html", p=None, lines=[],
                                   suppliers=suppliers, methods=methods,
                                   today=today_iso(), next_url=_safe_next())
        pid = next_id("PUR", "purchases")
        amount, _kg = _save_purchase(pid, header, lines, u["username"], True)
        audit(u["username"], "add", "purchase", pid,
              f"{header['supplier_id']} {amount}")
        flash(f"Purchase save ho gayi ({pid}) — {fmt_rs(amount)}", "ok")
        return redirect(_safe_next() or url_for("purchases"))
    return render_template("purchase_form.html", p=None, lines=[],
                           suppliers=suppliers, methods=methods,
                           today=today_iso(),
                           preselect=request.args.get("supplier_id", ""),
                           next_url=_safe_next())


@app.route("/purchases/<pid>/edit", methods=["GET", "POST"])
@owner_required
def purchase_edit(pid):
    u = current_user()
    con = _connect()
    try:
        p = con.execute("SELECT * FROM purchases WHERE id=?",
                        (pid,)).fetchone()
        suppliers = con.execute(
            "SELECT * FROM suppliers WHERE active=1 ORDER BY name").fetchall()
    finally:
        con.close()
    if not p or p["status"] == "voided":
        flash("Ye purchase edit nahi ho sakti (voided/missing)", "err")
        return redirect(url_for("purchases"))
    methods = list_methods()
    if request.method == "POST":
        header, lines = _parse_purchase_form()
        if not header["supplier_id"] or not lines:
            flash("Party aur kam az kam ek line zaroori hai", "err")
            return render_template("purchase_form.html", p=p,
                                   lines=_lines_of(pid), suppliers=suppliers,
                                   methods=methods, today=today_iso(),
                                   next_url=_safe_next())
        amount, _kg = _save_purchase(pid, header, lines, u["username"], False)
        audit(u["username"], "edit", "purchase", pid, f"{amount}")
        flash(f"Purchase update ho gayi ({pid}) — {fmt_rs(amount)}", "ok")
        return redirect(_safe_next() or url_for("purchases"))
    return render_template("purchase_form.html", p=p, lines=_lines_of(pid),
                           suppliers=suppliers, methods=methods,
                           today=today_iso(), next_url=_safe_next())


@app.route("/purchases/<pid>/void", methods=["GET", "POST"])
@owner_required
def purchase_void(pid):
    u = current_user()
    con = _connect()
    try:
        p = con.execute("SELECT * FROM purchases WHERE id=?",
                        (pid,)).fetchone()
    finally:
        con.close()
    if not p or p["status"] != "active":
        abort(404)
    if request.method == "POST":
        reason = (request.form.get("reason") or "").strip()
        if not reason:
            flash("Wajah likhna zaroori hai (audit trail)", "err")
            return render_template("void_form.html", title="Purchase Void",
                                   ref=pid)
        clear_purchase_stock(pid)
        con = _connect()
        try:
            con.execute(
                "UPDATE purchases SET status='voided', void_reason=?"
                " WHERE id=?", (reason, pid))
            con.execute(
                "UPDATE payments SET status='voided' WHERE purchase_id=?"
                " AND kind='purchase_payment'", (pid,))
            con.commit()
        finally:
            con.close()
        audit(u["username"], "void", "purchase", pid, reason)
        flash(f"{pid} void kar diya — stock bhi wapas ho gaya", "ok")
        return redirect(url_for("purchases"))
    return render_template("void_form.html", title="Purchase Void", ref=pid)


@app.route("/purchases/<pid>/delete", methods=["POST"])
@owner_required
def purchase_delete(pid):
    u = current_user()
    clear_purchase_stock(pid)
    con = _connect()
    try:
        p = con.execute("SELECT * FROM purchases WHERE id=?",
                        (pid,)).fetchone()
        if not p:
            con.close()
            abort(404)
        con.execute("DELETE FROM purchase_lines WHERE purchase_id=?", (pid,))
        con.execute(
            "DELETE FROM payments WHERE purchase_id=? AND"
            " kind='purchase_payment'", (pid,))
        con.execute(
            "UPDATE payments SET purchase_id='' WHERE purchase_id=?", (pid,))
        con.execute("DELETE FROM purchases WHERE id=?", (pid,))
        con.commit()
    finally:
        con.close()
    audit(u["username"], "delete", "purchase", pid,
          f"{p['supplier_id']} {p['amount']}")
    flash(f"Purchase {pid} delete ho gayi", "ok")
    return redirect(url_for("purchases"))


# ---------- payments v2 (owner) ----------

def _applies_to_from_form(kind):
    """Resolve the Previous/Current bucket for a payment form post:
    advances always land on Current, a payment tied to a purchase is
    forced to Current, otherwise the user's select wins (default
    Previous)."""
    purchase_id = request.form.get("purchase_id") or ""
    if kind == "advance" or purchase_id:
        return "current"
    v = (request.form.get("applies_to") or "previous").strip()
    return v if v in ("previous", "current") else "previous"


def _applies_label(r):
    """Display label for a payment row's bucket."""
    if r["kind"] in ("advance", "purchase_payment"):
        return "Current"
    return "Current" if (r["applies_to"] or "previous") == "current" \
        else "Previous"


@app.route("/payments")
@owner_required
def payments():
    con = _connect()
    try:
        rows = con.execute(
            "SELECT p.*, s.name AS sname FROM payments p"
            " JOIN suppliers s ON s.id=p.supplier_id"
            " ORDER BY COALESCE(p.date,'9999') DESC, p.id DESC").fetchall()
    finally:
        con.close()
    return render_template("payments.html", rows=rows)


def _payment_form_data():
    con = _connect()
    try:
        suppliers = con.execute(
            "SELECT * FROM suppliers WHERE active=1 ORDER BY name").fetchall()
        purs = con.execute(
            "SELECT p.id, p.date, p.amount, p.cash_paid, s.name AS sname"
            " FROM purchases p JOIN suppliers s ON s.id=p.supplier_id"
            " WHERE p.status='active' ORDER BY p.date DESC, p.id DESC"
        ).fetchall()
    finally:
        con.close()
    return suppliers, purs, list_methods()


@app.route("/payments/add", methods=["GET", "POST"])
@owner_required
def payment_add():
    u = current_user()
    suppliers, purs, methods = _payment_form_data()
    if request.method == "POST":
        sid = request.form.get("supplier_id") or ""
        amt = fnum(request.form.get("amount"))
        kind = request.form.get("kind") or "payment"
        if kind not in ("payment", "advance"):
            kind = "payment"
        mid = int(fnum(request.form.get("method_id")))
        mrow = method_by_id(mid) if mid else None
        if not sid or amt <= 0:
            flash("Party aur raqam (amount) zaroori hai", "err")
            return render_template("payment_form.html", p=None,
                                   suppliers=suppliers, purchases=purs,
                                   methods=methods, today=today_iso(),
                                   next_url=_safe_next())
        pid = next_id("PAY", "payments")
        con = _connect()
        try:
            con.execute(
                """INSERT INTO payments (id, date, supplier_id, amount,
                    status, cashbook_status, notes, created_by, created_at,
                    discount, method_id, method_name, kind, purchase_id,
                    applies_to)
                   VALUES (?,?,?, ?, 'active', 'na', ?, ?, ?, ?, ?, ?, ?,
                           ?, ?)""",
                (pid, request.form.get("date") or today_iso(), sid, amt,
                 (request.form.get("notes") or "").strip(), u["username"],
                 now_iso(), fnum(request.form.get("discount")),
                 mid, mrow["name"] if mrow else "",
                 kind, request.form.get("purchase_id") or "",
                 _applies_to_from_form(kind)),
            )
            con.commit()
        finally:
            con.close()
        audit(u["username"], "add", "payment", pid, f"{kind} {amt}")
        flash(f"Payment save ho gayi ({pid}) — {fmt_rs(amt)}", "ok")
        return redirect(_safe_next() or url_for("payments"))
    return render_template(
        "payment_form.html", p=None, suppliers=suppliers, purchases=purs,
        methods=methods, today=today_iso(),
        preselect=request.args.get("supplier_id", ""),
        pre_purchase=request.args.get("purchase_id", ""),
        next_url=_safe_next())


@app.route("/payments/<pid>/edit", methods=["GET", "POST"])
@owner_required
def payment_edit(pid):
    u = current_user()
    con = _connect()
    try:
        p = con.execute("SELECT * FROM payments WHERE id=?",
                        (pid,)).fetchone()
    finally:
        con.close()
    if not p or p["kind"] == "purchase_payment" or p["status"] == "voided":
        flash("Ye payment yahan se edit nahi hoti", "err")
        return redirect(url_for("payments"))
    suppliers, purs, methods = _payment_form_data()
    if request.method == "POST":
        amt = fnum(request.form.get("amount"))
        mid = int(fnum(request.form.get("method_id")))
        mrow = method_by_id(mid) if mid else None
        if amt <= 0:
            flash("Raqam zaroori hai", "err")
            return render_template("payment_form.html", p=p,
                                   suppliers=suppliers, purchases=purs,
                                   methods=methods, today=today_iso(),
                                   next_url=_safe_next())
        kind = request.form.get("kind") or "payment"
        if kind not in ("payment", "advance"):
            kind = "payment"
        con = _connect()
        try:
            con.execute(
                """UPDATE payments SET date=?, supplier_id=?, amount=?,
                    discount=?, method_id=?, method_name=?, kind=?,
                    purchase_id=?, applies_to=?, notes=? WHERE id=?""",
                (request.form.get("date") or p["date"],
                 request.form.get("supplier_id") or p["supplier_id"], amt,
                 fnum(request.form.get("discount")), mid,
                 mrow["name"] if mrow else "", kind,
                 request.form.get("purchase_id") or "",
                 _applies_to_from_form(kind),
                 (request.form.get("notes") or "").strip(), pid),
            )
            con.commit()
        finally:
            con.close()
        audit(u["username"], "edit", "payment", pid, f"{amt}")
        flash("Payment update ho gayi", "ok")
        return redirect(_safe_next() or url_for("payments"))
    return render_template("payment_form.html", p=p, suppliers=suppliers,
                           purchases=purs, methods=methods,
                           today=today_iso(), next_url=_safe_next())


@app.route("/payments/<pid>/void", methods=["GET", "POST"])
@owner_required
def payment_void(pid):
    u = current_user()
    con = _connect()
    try:
        p = con.execute("SELECT * FROM payments WHERE id=?",
                        (pid,)).fetchone()
    finally:
        con.close()
    if not p or p["status"] != "active" or p["kind"] == "purchase_payment":
        abort(404)
    if request.method == "POST":
        reason = (request.form.get("reason") or "").strip()
        if not reason:
            flash("Wajah likhna zaroori hai (audit trail)", "err")
            return render_template("void_form.html", title="Payment Void",
                                   ref=pid)
        con = _connect()
        try:
            con.execute(
                "UPDATE payments SET status='voided', void_reason=?"
                " WHERE id=?", (reason, pid))
            con.commit()
        finally:
            con.close()
        audit(u["username"], "void", "payment", pid, reason)
        flash(f"{pid} void kar diya", "ok")
        return redirect(url_for("payments"))
    return render_template("void_form.html", title="Payment Void", ref=pid)


@app.route("/payments/<pid>/delete", methods=["POST"])
@owner_required
def payment_delete(pid):
    u = current_user()
    con = _connect()
    try:
        p = con.execute("SELECT * FROM payments WHERE id=?",
                        (pid,)).fetchone()
        if not p or p["kind"] == "purchase_payment":
            con.close()
            abort(404)
        con.execute("DELETE FROM payments WHERE id=?", (pid,))
        con.commit()
    finally:
        con.close()
    audit(u["username"], "delete", "payment", pid, f"{p['amount']}")
    flash(f"Payment {pid} delete ho gayi", "ok")
    return redirect(url_for("payments"))


# ---------- payment methods / settings (owner) ----------

@app.route("/methods", methods=["GET", "POST"])
@owner_required
def methods():
    u = current_user()
    if request.method == "POST":
        name = (request.form.get("name") or "").strip()
        if not name:
            flash("Method ka naam zaroori hai", "err")
        else:
            con = _connect()
            try:
                con.execute(
                    "INSERT INTO payment_methods (name, active, sort,"
                    " created_at) VALUES (?, 1, 99, ?)", (name, now_iso()))
                con.commit()
                ok = True
            except Exception:
                ok = False
            finally:
                con.close()
            if ok:
                audit(u["username"], "add", "payment_method", name, "")
                flash(f"Method add ho gaya: {name}", "ok")
            else:
                flash("Ye method pehle se mojood hai", "err")
        return redirect(url_for("methods"))
    con = _connect()
    try:
        rows = con.execute(
            "SELECT * FROM payment_methods ORDER BY sort, name").fetchall()
    finally:
        con.close()
    return render_template("methods.html", rows=rows)


@app.route("/methods/<int:mid>/edit", methods=["POST"])
@owner_required
def method_edit(mid):
    u = current_user()
    name = (request.form.get("name") or "").strip()
    if not name:
        flash("Naam zaroori hai", "err")
        return redirect(url_for("methods"))
    con = _connect()
    try:
        con.execute(
            "UPDATE payment_methods SET name=?, active=? WHERE id=?",
            (name, 1 if request.form.get("active") else 0, mid))
        con.commit()
    except Exception:
        flash("Ye naam pehle se mojood hai", "err")
        return redirect(url_for("methods"))
    finally:
        con.close()
    audit(u["username"], "edit", "payment_method", str(mid), name)
    flash("Method update ho gaya", "ok")
    return redirect(url_for("methods"))


@app.route("/methods/<int:mid>/delete", methods=["POST"])
@owner_required
def method_delete(mid):
    u = current_user()
    con = _connect()
    try:
        m = method_by_id(mid)
        if not m:
            con.close()
            abort(404)
        n = con.execute(
            "SELECT COUNT(*) FROM payments WHERE method_id=?", (mid,)
        ).fetchone()[0] + con.execute(
            "SELECT COUNT(*) FROM purchases WHERE cash_method_id=?", (mid,)
        ).fetchone()[0]
        if n:
            con.close()
            flash("Ye method payments mein istemal hua hai — delete nahi"
                  " hoga, Edit se Off kar dein", "err")
            return redirect(url_for("methods"))
        con.execute("DELETE FROM payment_methods WHERE id=?", (mid,))
        con.commit()
    finally:
        con.close()
    audit(u["username"], "delete", "payment_method", str(mid), m["name"])
    flash(f"Method delete ho gaya: {m['name']}", "ok")
    return redirect(url_for("methods"))


# ---------- stock + usage ----------

@app.route("/stock")
@login_required
def stock():
    # v2.1: the list is ONLY items that were actually purchased
    # (recomputed from purchase lines — legacy materials never appear).
    # v2.2: canonical items (aliases merged), organised into groups.
    # v2.3: VIEW-ONLY page — all item-admin tools (group / rename /
    # merge) moved to /stock/manage so this page stays clean.
    # v2.4: SETS — body+blade cartons of a brand pair up into a brand
    # block (complete sets + extra) in Boxes & Packaging, and dhalai
    # models show their body/cup set counts.
    # v2.5: items explicitly assigned to a brand SECTION render inside
    # that brand's block (as extra lines, never counted in sets) and
    # not again in their material group; every item carries a bulk
    # date-wise detail (receipts + usage) for the expandable view.
    items = purchased_items()
    bc = db.dhalai_body_cup()
    det = db.stock_details()
    for it in items:
        d = bc.get(it["key"])
        if d:
            sets = min(d["body"], d["cup"])
            it["dhalai"] = {
                "body": d["body"], "cup": d["cup"], "sets": sets,
                "extra_body": max(0.0, d["body"] - d["cup"]),
                "extra_cup": max(0.0, d["cup"] - d["body"]),
            }
        it["detail"] = det.get(it["key"], {"receipts": [], "usage": []})
    groups = db.stock_groups(items)
    pairs, consumed = db.box_pairs(items)
    # v2.5: section-assigned items join their brand block as extras.
    pair_by_brand = {p["brand_key"]: p for p in pairs}
    for p in pairs:
        p["extras"] = []
    hidden = set(consumed)
    for it in items:
        bk = it.get("section") or ""
        if not bk or it["key"] in consumed:
            continue
        p = pair_by_brand.get(bk)
        if p is None:
            p = {"brand": db._brand_display(bk), "brand_key": bk,
                 "body": None, "blade": None, "sets": 0,
                 "extra_body": 0.0, "extra_blade": 0.0, "extras": []}
            pairs.append(p)
            pair_by_brand[bk] = p
        p["extras"].append(it)
        hidden.add(it["key"])
    if hidden:
        kept = []
        for g in groups:
            rows = [r for r in g["rows"] if r["key"] not in hidden]
            if g["name"] == "Boxes & Packaging":
                g["pairs"] = pairs
                g["count"] = len(rows) + sum(
                    (2 if p["body"] else 0) + len(p["extras"])
                    for p in pairs)
                g["rem_qty"] = sum(r["rem_qty"] for r in rows) + sum(
                    (p["body"]["rem_qty"] if p["body"] else 0)
                    + (p["blade"]["rem_qty"] if p["blade"] else 0)
                    + sum(e["rem_qty"] for e in p["extras"])
                    for p in pairs)
                g["rem_kg"] = sum(r["rem_kg"] for r in rows) + sum(
                    sum(e["rem_kg"] for e in p["extras"]) for p in pairs)
                g["rows"] = rows
                kept.append(g)
            elif rows:
                g["rows"] = rows
                g["count"] = len(rows)
                g["rem_qty"] = sum(r["rem_qty"] for r in rows)
                g["rem_kg"] = sum(r["rem_kg"] for r in rows)
                kept.append(g)
        if pairs and not any(g["name"] == "Boxes & Packaging"
                             for g in kept):
            kept.insert(0, {"name": "Boxes & Packaging", "rows": [],
                            "pairs": pairs,
                            "count": sum(
                                (2 if p["body"] else 0) + len(p["extras"])
                                for p in pairs),
                            "rem_qty": sum(
                                (p["body"]["rem_qty"] if p["body"] else 0)
                                + (p["blade"]["rem_qty"] if p["blade"]
                                   else 0)
                                + sum(e["rem_qty"] for e in p["extras"])
                                for p in pairs),
                            "rem_kg": sum(
                                sum(e["rem_kg"] for e in p["extras"])
                                for p in pairs)})
        groups = kept
    elif pairs:
        # no pairs were consumed and nothing is section-assigned —
        # cannot happen (pairs always consume their components), kept
        # for safety so brand blocks never silently vanish.
        for g in groups:
            if g["name"] == "Boxes & Packaging":
                g["pairs"] = pairs
    return render_template("stock.html", rows=items, groups=groups)


@app.route("/stock/manage")
@owner_required
def stock_manage():
    """Owner-only item administration — LIST screen (v3.0 forms
    style): a searchable list of clean item rows; tapping one opens
    its own page (stock_manage_item) with the group/section move,
    rename, merge and split forms as separate cards. The POST
    endpoints are unchanged since v2.3/v2.5."""
    items = purchased_items()
    groups = db.stock_groups(items)
    pairs, consumed = db.box_pairs(items)
    pair_by_brand = {p["brand_key"]: p for p in pairs}
    for p in pairs:
        p["members"] = [p["body"], p["blade"]]
    for it in items:
        bk = it.get("section") or ""
        if bk and it["key"] not in consumed:
            p = pair_by_brand.get(bk)
            if p is None:
                p = {"brand": db._brand_display(bk), "brand_key": bk,
                     "body": None, "blade": None, "sets": 0,
                     "extra_body": 0.0, "extra_blade": 0.0, "members": []}
                pairs.append(p)
                pair_by_brand[bk] = p
            p["members"].append(it)
    return render_template("stock_manage.html", groups=groups,
                           pairs=pairs,
                           brands=db.brand_sections(items),
                           aliases=db.list_aliases(items=items))


@app.route("/stock/manage/item")
@owner_required
def stock_manage_item():
    """One item's manage page (v3.0 forms style): header card, then
    one card per job — group/section move, rename, split, merge —
    plus its date-wise history. Display only; the POST endpoints
    (stock_move_group / stock_rename / stock_split / stock_merge)
    and their field names are exactly as before."""
    key = (request.args.get("key") or "").strip()
    it = purchased_item(key) if key else None
    if not it:
        flash("Item nahi mila", "err")
        return redirect(url_for("stock_manage"))
    items = purchased_items()
    groups = db.stock_groups(items)
    group_names = list(db.GROUP_ORDER) + sorted(
        {g["name"] for g in groups} - set(db.GROUP_ORDER))
    d = db.dhalai_body_cup().get(it["key"])
    if d:
        sets = min(d["body"], d["cup"])
        it["dhalai"] = {
            "body": d["body"], "cup": d["cup"], "sets": sets,
            "extra_body": max(0.0, d["body"] - d["cup"]),
            "extra_cup": max(0.0, d["cup"] - d["body"]),
        }
    it["detail"] = db.stock_details().get(
        it["key"], {"receipts": [], "usage": []})
    return render_template("stock_manage_item.html", x=it,
                           all_items=items, group_names=group_names,
                           brands=db.brand_sections(items),
                           split_modes=db.split_modes())


@app.route("/stock/merge", methods=["POST"])
@owner_required
def stock_merge():
    """Merge one canonical stock item into another (display only —
    purchase lines stay as they are; reversible via Unmerge)."""
    u = current_user()
    src = (request.form.get("variant_key") or "").strip()
    dst = (request.form.get("target_key") or "").strip()
    if not src or not dst or src == dst:
        flash("Merge ke liye do alag items chunein", "err")
        return redirect(url_for("stock"))
    target = purchased_item(dst)
    if not db.merge_items(src, dst):
        flash("Item nahi mila — merge nahi hua", "err")
        return redirect(url_for("stock"))
    audit(u["username"], "merge", "stock_item", src, f"-> {dst}")
    flash(f"Merge ho gaya: ab ye item '{target['name'] if target else dst}'"
          " ke saath ginna jayega", "ok")
    return redirect(url_for("stock"))


@app.route("/stock/rename", methods=["POST"])
@owner_required
def stock_rename():
    u = current_user()
    ck = (request.form.get("canonical_key") or "").strip()
    label = (request.form.get("label") or "").strip()
    if not ck or not purchased_item(ck):
        flash("Item nahi mila", "err")
        return redirect(url_for("stock"))
    db.set_label(ck, label)
    audit(u["username"], "rename", "stock_item", ck, label)
    flash("Item ka display naam update ho gaya", "ok")
    return redirect(url_for("stock"))


@app.route("/stock/aliases/delete", methods=["POST"])
@owner_required
def stock_alias_delete():
    u = current_user()
    variant = (request.form.get("variant") or "").strip()
    db.delete_alias(variant)
    audit(u["username"], "unmerge", "stock_item", variant, "")
    flash(f"Unmerge ho gaya: '{variant}' ab alag item hai", "ok")
    return redirect(url_for("stock"))


@app.route("/stock/move-group", methods=["POST"])
@app.route("/stock/move", methods=["POST"])
@owner_required
def stock_move_group():
    """Move an item to a material group OR a brand section (v2.5).
    Destination sources, in priority order: the free-text 'new_group'
    box (an exact brand naam/section match becomes a section, anything
    else a new/existing group), else the 'group_name'/'dest' select
    ('section:<brand_key>' values assign a section; picking a group
    clears any section). Drag-and-drop posts dest the same way."""
    u = current_user()
    ck = (request.form.get("canonical_key") or "").strip()
    it = purchased_item(ck) if ck else None
    if not it:
        flash("Item nahi mila", "err")
        return redirect(url_for("stock_manage"))
    items = purchased_items()
    brands = db.brand_sections(items)
    free = (request.form.get("new_group") or "").strip()
    sel = (request.form.get("group_name") or "").strip() \
        or (request.form.get("dest") or "").strip()
    if free:
        fk = db._norm_name(free)
        hit = None
        for bk, disp in brands.items():
            if fk == bk or fk == db._norm_name(disp):
                hit = bk
                break
        if hit:
            db.move_item(ck, "section", hit)
            audit(u["username"], "section", "stock_item", ck, hit)
            flash(f"Item '{brands[hit]}' section mein chala gaya", "ok")
        else:
            db.move_item(ck, "group", free)
            audit(u["username"], "group", "stock_item", ck, free)
            flash(f"Item '{free}' group mein chala gaya", "ok")
    elif sel.startswith("section:"):
        bk = sel.split(":", 1)[1]
        if bk not in brands:
            flash("Section nahi mila", "err")
            return redirect(url_for("stock_manage"))
        db.move_item(ck, "section", bk)
        audit(u["username"], "section", "stock_item", ck, bk)
        flash(f"Item '{brands[bk]}' section mein chala gaya", "ok")
    elif sel:
        db.move_item(ck, "group", sel)
        audit(u["username"], "group", "stock_item", ck, sel)
        flash(f"Item '{sel}' group mein chala gaya", "ok")
    else:
        flash("Item ya group nahi mila", "err")
    return redirect(url_for("stock_manage"))


@app.route("/stock/split", methods=["POST"])
@owner_required
def stock_split():
    """Split one stock item into a NEW item (owner tool, v2.5):
    qty/kg/body/cup per the item's kind; the split is an auditable
    movement pair — purchase history kabhi change nahi hoti."""
    u = current_user()
    ck = (request.form.get("canonical_key") or "").strip()
    ok, res = db.split_item(
        ck, request.form.get("dst_name"),
        qty=fnum(request.form.get("qty")),
        kg=fnum(request.form.get("kg")),
        body_qty=fnum(request.form.get("body_qty")),
        cup_qty=fnum(request.form.get("cup_qty")),
        actor=u["username"])
    if ok:
        flash(f"Split ho gaya: naya item '{request.form.get('dst_name')}'"
              " ban gaya", "ok")
    else:
        flash(res, "err")
    return redirect(url_for("stock_manage"))


@app.route("/stock/<path:item>")
@login_required
def material_ledger(item):
    it = purchased_item(item)
    if not it:
        # legacy /stock/<MAT-id> links: hop to the name-based page when
        # that material is a purchased item, else 404.
        con = _connect()
        try:
            m = con.execute("SELECT * FROM materials WHERE id=?",
                            (item,)).fetchone()
        finally:
            con.close()
        if m:
            resolved = purchased_item(m["name"])
            if resolved:
                return redirect(url_for("material_ledger",
                                        item=resolved["name"]))
        abort(404)
    con = _connect()
    try:
        # Canonical item: receipts cover EVERY variant line name, usage
        # moves cover every material in the group.
        names = it["variants"] or [it["name"]]
        ph = ",".join("?" * len(names))
        receipts = con.execute(
            "SELECT pl.*, p.date AS pdate, p.status AS pstatus,"
            " s.name AS sname FROM purchase_lines pl"
            " JOIN purchases p ON p.id = pl.purchase_id"
            " JOIN suppliers s ON s.id = p.supplier_id"
            f" WHERE p.status='active' AND TRIM(pl.name) IN ({ph})"
            " ORDER BY p.date DESC, p.id DESC, pl.sort, pl.id",
            names).fetchall()
        moves = []
        if it["material_ids"]:
            ph2 = ",".join("?" * len(it["material_ids"]))
            moves = con.execute(
                "SELECT l.*, m.name AS mname FROM stock_ledger l"
                " LEFT JOIN materials m ON m.id = l.material_id"
                f" WHERE l.material_id IN ({ph2})"
                " AND l.type IN ('OUT','ADJUST')"
                " ORDER BY l.date DESC, l.id DESC LIMIT 100",
                it["material_ids"]).fetchall()
    finally:
        con.close()
    return render_template("item_ledger.html", it=it, receipts=receipts,
                           moves=moves,
                           splits=db.splits_for(it["key"]))


@app.route("/usage/add", methods=["GET", "POST"])
@login_required
def usage_add():
    """Storekeeper's main screen: Item -> Qty -> Save. No money here.
    The item picker lists ONLY purchased items (v2.1)."""
    u = current_user()
    mats = [{"name": it["name"], "unit": it["unit"]}
            for it in purchased_items()]
    if request.method == "POST":
        item_name = (request.form.get("item_name") or "").strip()
        mid = request.form.get("material_id") or ""  # legacy posts
        qty = fnum(request.form.get("qty"))
        d = request.form.get("date") or today_iso()
        note = (request.form.get("notes") or "").strip()
        it = None
        if item_name:
            # v2.2: the picker lists canonical items (display labels);
            # the OUT row posts against the group's primary material and
            # counts against the COMBINED remaining of the whole group.
            it = purchased_item(item_name)
            if not it:
                flash("Item nahi mila", "err")
                return render_template("usage_form.html", mats=mats,
                                       today=today_iso())
            mid = it["material_id"] or get_or_create_material(
                it["variants"][0] if it["variants"] else it["name"],
                it["unit"])
        if not mid or qty <= 0:
            flash("Item aur qty zaroori hai", "err")
            return render_template("usage_form.html", mats=mats,
                                   today=today_iso())
        con = _connect()
        try:
            m = con.execute("SELECT * FROM materials WHERE id=?",
                            (mid,)).fetchone()
        finally:
            con.close()
        if not m:
            flash("Item nahi mila", "err")
            return render_template("usage_form.html", mats=mats,
                                   today=today_iso())
        try:
            _id, bal = post_stock(d, mid, "OUT", qty, "",
                                  u["username"], note or "daily usage")
        except ValueError as e:
            flash(str(e), "err")
            return render_template("usage_form.html", mats=mats,
                                   today=today_iso())
        audit(u["username"], "usage", "stock", mid, f"OUT {qty} bal={bal}")
        if it:
            fresh = purchased_item(it["key"])
            rem = fresh["rem_qty"] if it["unit"] == "pcs" \
                else fresh["rem_kg"]
            flash(f"Save ho gaya: {it['name']} — {fmt_qty(qty)}"
                  f" {it['unit']} use. Bacha: {fmt_qty(rem)}", "ok")
        else:
            flash(f"Save ho gaya: {m['name']} — {fmt_qty(qty)}"
                  f" {m['unit']} use. Bacha: {fmt_qty(bal)}", "ok")
        return redirect(url_for("stock"))
    return render_template("usage_form.html", mats=mats, today=today_iso())


@app.route("/stock/movement/<int:lid>/delete", methods=["POST"])
@owner_required
def movement_delete(lid):
    u = current_user()
    con = _connect()
    try:
        mv = con.execute("SELECT * FROM stock_ledger WHERE id=?",
                         (lid,)).fetchone()
        if not mv:
            con.close()
            abort(404)
        con.execute("DELETE FROM stock_ledger WHERE id=?", (lid,))
        con.commit()
    finally:
        con.close()
    audit(u["username"], "delete", "stock_movement", str(lid),
          f"{mv['type']} {mv['qty']}")
    flash("Movement delete ho gayi", "ok")
    con = _connect()
    try:
        mrow = con.execute("SELECT name FROM materials WHERE id=?",
                           (mv["material_id"],)).fetchone()
    finally:
        con.close()
    if mrow and purchased_item(mrow["name"]):
        return redirect(url_for("material_ledger", item=mrow["name"]))
    return redirect(url_for("stock"))


@app.route("/stock/adjust", methods=["GET", "POST"])
@owner_required
def stock_adjust():
    """Owner-only correction with reason (audit trail)."""
    u = current_user()
    con = _connect()
    try:
        mats = con.execute(
            "SELECT * FROM materials WHERE active=1 ORDER BY name").fetchall()
    finally:
        con.close()
    if request.method == "POST":
        mid = request.form.get("material_id") or ""
        qty = fnum(request.form.get("qty"))  # signed: + add, - subtract
        reason = (request.form.get("reason") or "").strip()
        d = request.form.get("date") or today_iso()
        if not mid or qty == 0 or not reason:
            flash("Item, qty (+/-) aur wajah — teeno zaroori hain", "err")
            return render_template("adjust_form.html", mats=mats,
                                   today=today_iso())
        try:
            _id, bal = post_stock(d, mid, "ADJUST", qty, "",
                                  u["username"], f"ADJUST: {reason}")
        except ValueError as e:
            flash(str(e), "err")
            return render_template("adjust_form.html", mats=mats,
                                   today=today_iso())
        audit(u["username"], "adjust", "stock", mid, f"{qty}: {reason}")
        flash(f"Adjustment save: bachat ab {fmt_qty(bal)}", "ok")
        return redirect(url_for("material_ledger", mid=mid))
    return render_template("adjust_form.html", mats=mats, today=today_iso())


@app.route("/materials/add", methods=["POST"])
@owner_required
def material_add():
    u = current_user()
    name = (request.form.get("name") or "").strip()
    unit = (request.form.get("unit") or "kg").strip() or "kg"
    if not name:
        flash("Item ka naam zaroori hai", "err")
        return redirect(url_for("stock"))
    mid = get_or_create_material(name, unit)
    con = _connect()
    try:
        con.execute(
            "UPDATE materials SET opening_stock=?, min_level=?, unit=?"
            " WHERE id=?",
            (fnum(request.form.get("opening_stock")),
             fnum(request.form.get("min_level")), unit, mid),
        )
        con.commit()
    finally:
        con.close()
    audit(u["username"], "add", "material", mid, name)
    flash(f"Item add/update: {name}", "ok")
    return redirect(url_for("stock"))


# ---------- integrations (owner) — CashBook OFF ----------

@app.route("/integrations")
@owner_required
def integrations():
    con = _connect()
    try:
        rows = con.execute(
            "SELECT * FROM cashbook_sync ORDER BY id DESC LIMIT 50").fetchall()
    finally:
        con.close()
    pend, fail = cashbook_link.queue_counts()
    return render_template("integrations.html", rows=rows,
                           pending=pend, failed=fail, direct=False)


# ---------- excel exports (owner) ----------

def _style_xlsx(ws, headers):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    blue = PatternFill("solid", fgColor="047857")
    wf = Font(color="FFFFFF", bold=True, size=11)
    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = blue
        cell.font = wf
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    ws.row_dimensions[1].height = 26
    for i, h in enumerate(headers, 1):
        letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letter].width = max(14, min(36, len(str(h)) + 8))
    return border


def _finish_xlsx(ws, filename, headers):
    from openpyxl.styles import Alignment
    border = _style_xlsx(ws, headers)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row,
                            max_col=len(headers)):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center")
    path = f"/tmp/{filename}"
    ws.parent.save(path)
    return send_file(path, as_attachment=True, download_name=filename)


def _finish_xlsx_pro(ws, filename, headers, title, context,
                     money_cols=(), total_pred=None, total_last=False,
                     widths=None, landscape=False, download_name=None):
    """Professional statement finish (v3.3) for the party-ledger and
    bakaya sheets. The sheet must arrive with the ORIGINAL header texts
    in row 1 and the original data rows below; this inserts the title
    block (row 1 merged title, row 2 context, row 3 blank) so the header
    lands on row 4 and the data is unchanged below it. Amounts become
    '#,##0' numbers, totals rows bold on a light green fill, panes
    frozen under the header."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    ncol = len(headers)
    ws.insert_rows(1, 3)
    green_fill = PatternFill("solid", fgColor="2E7D32")
    light_fill = PatternFill("solid", fgColor="E8F3E8")
    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    tcell = ws.cell(row=1, column=1, value=title)
    tcell.font = Font(bold=True, size=14, color="2E7D32")
    tcell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 24
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
    ccell = ws.cell(row=2, column=1, value=context)
    ccell.font = Font(size=10, color="5B6B5B")
    ccell.alignment = Alignment(vertical="center")
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=c)
        cell.fill = green_fill
        cell.font = Font(color="FFFFFF", bold=True, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    ws.row_dimensions[4].height = 26
    for r in range(5, ws.max_row + 1):
        vals = tuple(ws.cell(row=r, column=c).value
                     for c in range(1, ncol + 1))
        is_total = bool((total_pred and total_pred(vals))
                        or (total_last and r == ws.max_row))
        for c in range(1, ncol + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            cell.alignment = Alignment(
                vertical="center",
                horizontal="right" if c in money_cols else "left")
            if c in money_cols and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"
            if is_total:
                cell.font = Font(bold=True)
                cell.fill = light_fill
    for i, h in enumerate(headers, 1):
        letter = ws.cell(row=4, column=i).column_letter
        w = (widths or {}).get(i) or max(14, min(36, len(str(h)) + 8))
        ws.column_dimensions[letter].width = w
    ws.freeze_panes = "A5"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    if landscape:
        ws.page_setup.orientation = "landscape"
    path = f"/tmp/{filename}"
    ws.parent.save(path)
    return send_file(path, as_attachment=True,
                     download_name=download_name or filename)


@app.route("/export/<name>.xlsx")
@owner_required
def export_xlsx(name):
    import openpyxl
    con = _connect()
    try:
        if name == "suppliers":
            rows = con.execute(
                "SELECT * FROM suppliers ORDER BY name").fetchall()
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Parties"
            H = ["ID", "Name", "Category", "Phone", "Opening Balance",
                 "Previous Remaining (Rs)", "Current Bakaya (Rs)",
                 "Outstanding (Rs)", "Active"]
            ws.append(H)
            bals = all_party_balances(
                con, {r["id"]: (r["opening_balance"] or 0) for r in rows})
            for r in rows:
                bal = bals[r["id"]]
                ws.append([r["id"], r["name"], r["category"], r["phone"],
                           r["opening_balance"] or 0,
                           round(bal["prev_remaining"], 0),
                           round(bal["current_bakaya"], 0),
                           round(bal["total"], 0),
                           "Yes" if r["active"] else "No"])
            return _finish_xlsx(ws, "suppliers.xlsx", H)
        if name == "purchases":
            rows = con.execute(
                "SELECT p.*, s.name AS sname FROM purchases p"
                " JOIN suppliers s ON s.id=p.supplier_id"
                " ORDER BY p.date DESC, p.id DESC").fetchall()
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Purchases"
            H = ["ID", "Date", "Party", "Mode", "Lines", "Total Kg",
                 "Rate/Kg", "Amount (Rs)", "Cash Paid", "Cash Method",
                 "Bakaya (Rs)", "Status", "Note"]
            ws.append(H)
            lines_map = _bulk_lines(con, [r["id"] for r in rows])
            linked = _bulk_linked(
                con, [r["id"] for r in rows if r["status"] == "active"])
            for r in rows:
                lines = lines_map.get(r["id"], [])
                bakaya = ((r["amount"] or 0) - (r["cash_paid"] or 0)
                          - linked.get(r["id"], 0)) \
                    if r["status"] == "active" else 0
                ws.append([r["id"], r["date"], r["sname"], r["mode"],
                           lines_summary(lines), r["total_kg"],
                           r["rate_per_kg"], r["amount"], r["cash_paid"],
                           r["cash_method_name"],
                           round(bakaya, 0),
                           r["status"], r["notes"] or ""])
            return _finish_xlsx(ws, "purchases.xlsx", H)
        if name == "payments":
            rows = con.execute(
                "SELECT p.*, s.name AS sname FROM payments p"
                " JOIN suppliers s ON s.id=p.supplier_id"
                " ORDER BY COALESCE(p.date,'9999') DESC, p.id DESC"
            ).fetchall()
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Payments"
            H = ["ID", "Date", "Party", "Kind", "Applies To",
                 "Amount (Rs)", "Discount (Rs)", "Total (Rs)", "Method",
                 "Purchase", "Status", "Notes"]
            ws.append(H)
            for r in rows:
                ws.append([r["id"], r["date"] or "", r["sname"], r["kind"],
                           _applies_label(r),
                           r["amount"], r["discount"] or 0,
                           (r["amount"] or 0) + (r["discount"] or 0),
                           r["method_name"] or r["paid_from"] or "",
                           r["purchase_id"] or "", r["status"],
                           r["notes"] or ""])
            return _finish_xlsx(ws, "payments.xlsx", H)
        if name == "previous-bakaya":
            # v3.2: the Previous Balances Form as a sheet — one row per
            # party that HAS a previous (opening) balance.
            rows = con.execute(
                "SELECT * FROM suppliers WHERE active=1 ORDER BY name"
            ).fetchall()
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Previous Bakaya"
            H = ["Party", "Opening (Rs)", "Payments Against Previous (Rs)",
                 "Remaining (Rs)"]
            ws.append(H)
            bals = all_party_balances(
                con, {r["id"]: (r["opening_balance"] or 0) for r in rows})
            t_open = t_paid = t_rem = 0.0
            for r in rows:
                if (r["opening_balance"] or 0) <= 0:
                    continue
                bal = bals[r["id"]]
                t_open += round(bal["opening"], 0)
                t_paid += round(bal["prev_paid"], 0)
                t_rem += round(bal["prev_remaining"], 0)
                ws.append([r["name"], round(bal["opening"], 0),
                           round(bal["prev_paid"], 0),
                           round(bal["prev_remaining"], 0)])
            ws.append(["Total", t_open, t_paid, t_rem])
            return _finish_xlsx_pro(
                ws, "previous-bakaya.xlsx", H,
                "ZK FAN — Previous Bakaya",
                f"Saari parties · Generated: {date.today().isoformat()}",
                money_cols=(2, 3, 4),
                total_pred=lambda vals: vals[0] == "Total")
        if name == "current-bakaya":
            # v3.2: every bakaya-carrying purchase, grouped by party
            # with a subtotal row under each party.
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Current Bakaya"
            H = ["Date", "Party", "Purchase No", "Total (Rs)",
                 "Cash Diya (Rs)", "Bakaya Bacha (Rs)"]
            ws.append(H)
            brows = _bakaya_purchase_rows(con)
            by_party = {}
            for x in brows:
                by_party.setdefault(x["row"]["sname"], []).append(x)
            ordered = sorted(
                by_party.items(),
                key=lambda kv: -sum(x["bakaya"] for x in kv[1]))
            for sname, xs in ordered:
                for x in xs:
                    p = x["row"]
                    ws.append([p["date"], sname, p["id"],
                               round(p["amount"] or 0, 0),
                               round(x["cash_diya"], 0),
                               round(x["bakaya"], 0)])
                ws.append(["", sname, "Subtotal", "", "",
                           round(sum(x["bakaya"] for x in xs), 0)])
            return _finish_xlsx_pro(
                ws, "current-bakaya.xlsx", H,
                "ZK FAN — Current Bakaya",
                f"Saari parties · Generated: {date.today().isoformat()}",
                money_cols=(4, 5, 6),
                total_pred=lambda vals: len(vals) > 2
                and vals[2] == "Subtotal")
        if name == "stock":
            # v2.1: purchased items only — same source as the Stock page.
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Stock"
            H = ["Item", "Unit", "Received Qty", "Received Kg",
                 "Used Qty", "Used Kg", "Adjust Qty", "Adjust Kg",
                 "Remaining Qty", "Remaining Kg"]
            ws.append(H)
            for it in purchased_items():
                ws.append([it["name"], it["unit"],
                           round(it["in_qty"], 3), round(it["in_kg"], 3),
                           round(it["out_qty"], 3), round(it["out_kg"], 3),
                           round(it["adj_qty"], 3), round(it["adj_kg"], 3),
                           round(it["rem_qty"], 3), round(it["rem_kg"], 3)])
            return _finish_xlsx(ws, "stock.xlsx", H)
    finally:
        con.close()
    abort(404)


@app.route("/export/party-ledger/<sid>.xlsx")
@owner_required
def export_party_ledger(sid):
    import openpyxl
    con = _connect()
    try:
        s = con.execute("SELECT * FROM suppliers WHERE id=?",
                        (sid,)).fetchone()
    finally:
        con.close()
    if not s:
        abort(404)
    rows, bal = _ledger_rows(s)
    split = party_balances(sid)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ledger"
    H = ["Date", "Type", "Previous/Current", "Ref", "Detail",
         "Plus (Rs)", "Minus (Rs)", "Balance (Rs)"]
    ws.append(H)
    for r in rows:
        ws.append([r["date"], r["type"], r.get("tag", ""), r["ref"],
                   r["detail"], r["plus"] or "", r["minus"] or "",
                   r["bal"]])
    context = (f"{s['name']} · {s['id']} · Previous remaining: "
               f"Rs {split['prev_remaining']:,.0f} · Current remaining: "
               f"Rs {split['current_bakaya']:,.0f} · Total outstanding: "
               f"Rs {bal:,.0f} · Generated: {date.today().isoformat()}")
    return _finish_xlsx_pro(
        ws, f"ledger_{sid}.xlsx", H,
        "ZK FAN — Supplier Statement", context,
        money_cols=(6, 7, 8), total_last=True, widths={5: 52},
        landscape=True, download_name=f"ledger_{s['name']}.xlsx")


def _statement_pdf_bytes(s, rows, bal, split):
    """Professional A4 statement PDF for one party (reportlab Platypus,
    v3.3). Rows/order come straight from _ledger_rows — the same data
    as the on-screen ledger and the statement xlsx; no parallel math.
    Only called when HAVE_REPORTLAB is true."""
    import io
    from xml.sax.saxutils import escape
    buf = io.BytesIO()
    green = rl_colors.HexColor("#2e7d32")
    light = rl_colors.HexColor("#e8f3e8")
    grey = rl_colors.HexColor("#5b6b5b")
    grid = rl_colors.HexColor("#c9d4c9")
    void_grey = rl_colors.HexColor("#9a9a9a")

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(grey)
        canvas.drawString(15 * RL_MM, 10 * RL_MM,
                          "ZK FAN — Supplier & Stock System se tayyar")
        canvas.drawRightString(
            RL_A4[0] - 15 * RL_MM, 10 * RL_MM, f"Page {doc.page}")
        canvas.restoreState()

    doc = RL_Doc(buf, pagesize=RL_A4,
                 leftMargin=15 * RL_MM, rightMargin=15 * RL_MM,
                 topMargin=14 * RL_MM, bottomMargin=16 * RL_MM,
                 title=f"ZK FAN Statement — {s['name']}", author="ZK FAN")
    el = [RL_Para("ZK FAN", RL_PStyle(
              "mh", fontName="Helvetica-Bold", fontSize=20, leading=24,
              textColor=green)),
          RL_Para("Supplier Statement", RL_PStyle(
              "msub", fontName="Helvetica", fontSize=12, leading=15,
              textColor=grey)),
          RL_Para(f"Generated: {date.today().isoformat()}", RL_PStyle(
              "mgen", fontName="Helvetica", fontSize=9, leading=12,
              textColor=grey)),
          RL_Spacer(1, 4 * RL_MM),
          RL_Para(f"<b>{escape(str(s['name']))}</b> · "
                  f"{escape(str(s['id']))}", RL_PStyle(
                      "party", fontName="Helvetica", fontSize=11,
                      leading=14)),
          RL_Spacer(1, 3 * RL_MM)]
    band = RL_Table(
        [["Previous balance bacha", "Current bakaya", "Total outstanding"],
         [fmt_rs(split["prev_remaining"]), fmt_rs(split["current_bakaya"]),
          fmt_rs(bal)]],
        colWidths=[60 * RL_MM, 60 * RL_MM, 60 * RL_MM])
    band.setStyle(RL_TStyle([
        ("BACKGROUND", (0, 0), (-1, 0), light),
        ("TEXTCOLOR", (0, 0), (-1, 0), grey),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 8.5),
        ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 1), (-1, 1), 13),
        ("BOX", (0, 0), (-1, -1), 0.7, grid),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, grid),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 8)]))
    el.append(band)
    el.append(RL_Spacer(1, 4 * RL_MM))
    det = RL_PStyle("det", fontName="Helvetica", fontSize=8.5, leading=11)
    det_void = RL_PStyle("detv", parent=det, textColor=void_grey)
    data = [["Date", "Type", "Detail", "Plus (Rs)", "Minus (Rs)",
             "Balance (Rs)"]]
    void_rows = []
    for r in rows:
        typ = r["type"] + (f" ({r['tag']})" if r.get("tag") else "")
        if r["void"]:
            typ += " (void)"
            void_rows.append(len(data))
        data.append([
            str(r["date"]), typ,
            RL_Para(escape(str(r["detail"] or "")),
                    det_void if r["void"] else det),
            f"{r['plus']:,.0f}" if r["plus"] else "—",
            f"{r['minus']:,.0f}" if r["minus"] else "—",
            f"{r['bal']:,.0f}"])
    data.append(["Final outstanding", "", "", "", "", f"{bal:,.0f}"])
    tbl = RL_Table(data, repeatRows=1,
                   colWidths=[22 * RL_MM, 30 * RL_MM, 52 * RL_MM,
                              24 * RL_MM, 24 * RL_MM, 28 * RL_MM])
    style = [("BACKGROUND", (0, 0), (-1, 0), green),
             ("TEXTCOLOR", (0, 0), (-1, 0), rl_colors.white),
             ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
             ("FONTSIZE", (0, 0), (-1, 0), 9.5),
             ("FONTNAME", (0, 1), (-1, -2), "Helvetica"),
             ("FONTSIZE", (0, 1), (-1, -2), 9),
             ("ALIGN", (3, 0), (-1, -1), "RIGHT"),
             ("GRID", (0, 0), (-1, -1), 0.4, grid),
             ("VALIGN", (0, 0), (-1, -1), "TOP"),
             ("TOPPADDING", (0, 0), (-1, -1), 4),
             ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
             ("SPAN", (0, -1), (4, -1)),
             ("ALIGN", (0, -1), (0, -1), "RIGHT"),
             ("BACKGROUND", (0, -1), (-1, -1), light),
             ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
             ("FONTSIZE", (0, -1), (-1, -1), 9.5)]
    for i in void_rows:
        style.append(("TEXTCOLOR", (0, i), (-1, i), void_grey))
    tbl.setStyle(RL_TStyle(style))
    el.append(tbl)
    doc.build(el, onFirstPage=footer, onLaterPages=footer)
    buf.seek(0)
    return buf


@app.route("/export/party-ledger/<sid>.pdf")
@owner_required
def export_party_ledger_pdf(sid):
    """Server-side statement PDF (v3.3). Degrades gracefully when the
    server has no reportlab: friendly flash + back to the ledger (the
    page also hides its PDF button in that case)."""
    con = _connect()
    try:
        s = con.execute("SELECT * FROM suppliers WHERE id=?",
                        (sid,)).fetchone()
        if not s:
            abort(404)
        if not HAVE_REPORTLAB:
            flash("PDF export is machine par available nahi — Excel"
                  " wala button use karein", "err")
            return redirect(url_for("supplier_ledger", sid=sid))
        rows, bal = _ledger_rows(s, con)
        split = party_balances(sid, con)
    finally:
        con.close()
    buf = _statement_pdf_bytes(s, rows, bal, split)
    return send_file(buf, mimetype="application/pdf", as_attachment=False,
                     download_name=f"statement_{s['name']}.pdf")


# ---------- users & backup (owner) ----------

@app.route("/users", methods=["GET", "POST"])
@owner_required
def users():
    u = current_user()
    if request.method == "POST":
        un = (request.form.get("username") or "").strip()
        role = request.form.get("role") or "storekeeper"
        pw = request.form.get("password") or ""
        if not un or len(pw) < 6 or role not in ("owner", "storekeeper"):
            flash("Username, role aur 6+ character password zaroori hai",
                  "err")
        else:
            con = _connect()
            try:
                con.execute(
                    "INSERT INTO users (username, password_hash, role,"
                    " must_change_pw, active, created_at)"
                    " VALUES (?,?,?,1,1,?)",
                    (un, generate_password_hash(pw), role, now_iso()),
                )
                con.commit()
                ok = True
            except Exception:
                ok = False
            finally:
                con.close()
            if not ok:
                flash("Username pehle se mojood hai", "err")
                return redirect(url_for("users"))
            audit(u["username"], "add", "user", un, role)
            flash(f"User ban gaya: {un} ({role})", "ok")
        return redirect(url_for("users"))
    con = _connect()
    try:
        rows = con.execute(
            "SELECT id, username, role, must_change_pw, active, created_at"
            " FROM users ORDER BY id").fetchall()
    finally:
        con.close()
    return render_template("users.html", rows=rows)


@app.route("/backup")
@owner_required
def backup():
    u = current_user()
    bdir = os.path.join(db.BASE_DIR, "backups")
    os.makedirs(bdir, exist_ok=True)
    stamp = now_iso().replace(":", "").replace("-", "")[:14]
    dest = os.path.join(bdir, f"supplier_stock_{stamp}.db")
    try:
        shutil.copy2(db.DB_PATH, dest)
        files = sorted(f for f in os.listdir(bdir) if f.endswith(".db"))
        for old in files[:-30]:
            os.remove(os.path.join(bdir, old))
        audit(u["username"], "backup", "db", dest, "")
        flash(f"Backup ho gaya: {os.path.basename(dest)}", "ok")
    except Exception as e:
        flash(f"Backup fail: {e}", "err")
    return redirect(url_for("integrations"))


# ---------- error pages ----------

@app.errorhandler(403)
def e403(_e):
    return render_template("error.html",
                           msg="Aapko is page ki ijazat nahi hai."), 403


@app.errorhandler(404)
def e404(_e):
    return render_template("error.html", msg="Page nahi mila."), 404


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5051, debug=True)
